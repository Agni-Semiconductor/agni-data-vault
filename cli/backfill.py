#!/usr/bin/env python3
"""Sure-only backfill for the Agni Data Vault.

Walks a folder of Keithley Clarius Excel exports, writes labels only where
the evidence proves them (E1 file metadata, E2 filename tokens, E3 folder
tokens, E4 owner declaration), plans one sample + one measurement per
workbook, and queues everything unproven for human review.
"""
import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import sys
import time
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from cli.vault import ApiError, VaultClient  # noqa: E402

SAMPLE_ID = "HY_20nm_highT_0827"
EVIDENCE_OWNER = "Spencer 2026-09-09 chat; folder suffix HY"
KIND_BY_TOKEN = {"DC": "dciv", "AC": "aciv", "PUND": "pund", "CV": "cv"}
RUN_HDR = re.compile(r"^Run(\d+)$", re.I)
FNAME_RE = re.compile(r"^(\d+)-([A-Za-z]+)-(\d+)(?:[-_ ].*)?$")
SCREENSHOT_EXTS = {".bmp": "image/bmp", ".png": "image/png"}


def load_env(path=None):
    env = Path(path) if path else ROOT / ".env.local"
    if not env.exists():
        return
    for line in env.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip("\"'"))


def make_client():
    load_env()
    url, key = os.getenv("VAULT_API_URL"), os.getenv("VAULT_API_KEY")
    if not url or not key:
        sys.exit("VAULT_API_URL / VAULT_API_KEY missing (expected in .env.local)")
    return VaultClient(url, key)


def parse_filename(name):
    """E2: '<pad>-<KIND>-<rep>[-suffix].xlsx' -> {pad_dim_um, kind, replicate} or None."""
    path = Path(name)
    if path.suffix.lower() != ".xlsx":
        return None
    stem = path.stem
    while stem.lower().endswith(".xlsx"):
        stem = stem[:-5]
    match = FNAME_RE.match(stem)
    if not match:
        return None
    token = match.group(2).upper()
    return {
        "pad_dim_um": int(match.group(1)),
        "kind": KIND_BY_TOKEN.get(token),
        "replicate": int(match.group(3)),
        "token": token,
    }


def parse_folder(name):
    """E3: subfolder name -> (temperature_c, thermal_history)."""
    text = name.strip()
    thermal = None
    match = re.search(r"check[-_ ]?after\s*(\d+)\s*C", text, re.I)
    if match:
        thermal = "after " + match.group(1) + "C"
    temp = None
    match = re.search(r"(\d+(?:\.\d+)?)\s*C", text)
    if match:
        temp = int(float(match.group(1)))
    elif text.upper().startswith("RT"):
        temp = 20
    return temp, thermal


def kind_from_test_name(name):
    text = (name or "").lower()
    if "dc-iv" in text or "dc iv" in text or "dciv" in text:
        return "dciv"
    if "pund" in text:
        return "pund"
    if "c-v" in text or re.search(r"\bcv\b", text):
        return "cv"
    if "hysteresis" in text or re.search(r"\bac\b", text):
        return "aciv"
    return None


def kind_agreement(token_kind, test_names):
    """E1+E2: kind is written only when the Settings Test Name and the filename token agree."""
    kinds = {kind_from_test_name(n) for n in test_names if n}
    kinds.discard(None)
    if not kinds:
        return None, "no recognizable Test Name in Settings (filename token: %s)" % (token_kind or "none")
    if len(kinds) > 1:
        return None, "Settings Test Names disagree across runs: %s" % sorted(kinds)
    found = kinds.pop()
    if token_kind is None:
        return None, "filename has no kind token; Test Name says %s" % found
    if found != token_kind:
        return None, "Test Name says %s but filename token says %s" % (found, token_kind)
    return found, None


def resolve_kind(token_kind, test_names, module_names):
    """Kind from E1 (Test Name / Module Name) when it agrees with the E2 filename token;
    E2 alone when Settings names nothing; queued (None) on conflict or no evidence."""
    e1 = {kind_from_test_name(n) for n in list(test_names) + list(module_names)}
    e1.discard(None)
    if len(e1) > 1:
        return None, None, None, "Settings Test Name/Module Name imply several kinds: %s" % sorted(e1)
    if e1:
        found = e1.pop()
        if token_kind is None:
            return None, None, None, "filename has no kind token; Settings says %s" % found
        if found != token_kind:
            return None, None, None, "Test Name/Module Name says %s but filename token says %s" % (found, token_kind)
        return found, "E1", "Settings!Test Name/Module Name agrees with filename token", None
    if token_kind:
        return token_kind, "E2", "filename kind token", None
    return None, None, None, "no kind evidence in Settings or filename"


def _to_float(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if isinstance(value, str):
        match = re.match(r"^\s*(-?\d+(?:\.\d+)?)", value)
        if match:
            return float(match.group(1))
    return None


def _freq_to_khz(value):
    if value is None:
        return None
    match = re.match(r"^\s*([0-9.]+)\s*(khz|mhz|hz|k|m)?\.?\s*$", str(value).strip().lower())
    if not match:
        return None
    number = float(match.group(1))
    unit = match.group(2) or ""
    return {"hz": number / 1000.0, "khz": number, "k": number,
            "mhz": number * 1000.0, "m": number * 1000.0}[unit]


def _parse_executed(value):
    if isinstance(value, dt.datetime):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def parse_settings_rows(rows):
    """Parse vertical key/value Settings rows into per-run blocks (one per Run<N> header)."""
    runs, current = [], None
    for row in rows:
        first = str(row[0]).strip() if row and row[0] is not None else ""
        match = RUN_HDR.match(first)
        if match:
            current = {"run": int(match.group(1)), "test_name": None, "executed": None,
                       "clarius_version": None, "module_name": None, "keys": [],
                       "sweep_v": None, "sweep_src": None, "frequency_khz": None, "freq_src": None}
            runs.append(current)
            continue
        if current is None or not first:
            continue
        key = first.lower()
        value = row[1] if len(row) > 1 else None
        current["keys"].append(key)
        if key == "test name":
            current["test_name"] = str(value).strip() if value is not None else None
        elif key == "module name":
            current["module_name"] = str(value).strip() if value is not None else None
        elif key == "last executed":
            current["executed"] = _parse_executed(value)
        elif key == "clarius+ version":
            current["clarius_version"] = str(value).strip() if value is not None else None
        elif current["sweep_v"] is None and ("start" in key or "stop" in key or "bias" in key):
            number = _to_float(value)
            if number is not None:
                current["sweep_v"], current["sweep_src"] = number, first
        elif current["frequency_khz"] is None and "freq" in key:
            khz = _freq_to_khz(value)
            if khz is not None:
                current["frequency_khz"], current["freq_src"] = khz, first
    return runs


def measured_on_from_runs(runs):
    executed = [r["executed"] for r in runs if r.get("executed")]
    return min(executed).date().isoformat() if executed else None


def executed_at_list(runs):
    ordered = sorted((r for r in runs if r.get("executed")), key=lambda r: r["run"])
    return [r["executed"].isoformat() for r in ordered]


def read_workbook(path):
    """E1 read: Settings blocks, Run<N> sheet numbers, and only the first data-sheet header row."""
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        run_numbers = sorted(int(RUN_HDR.match(n).group(1)) for n in names if RUN_HDR.match(n))
        settings_rows = []
        if "Settings" in names:
            for row in workbook["Settings"].iter_rows(values_only=True):
                settings_rows.append(list(row[:3]))
        columns, first_run_sheet = [], None
        for name in names:
            if RUN_HDR.match(name):
                first_run_sheet = name
                break
        if first_run_sheet:
            for row in workbook[first_run_sheet].iter_rows(max_row=1, values_only=True):
                columns = [str(c).strip() for c in row if c is not None and str(c).strip()]
                break
    finally:
        workbook.close()
    return {"run_numbers": run_numbers, "runs": parse_settings_rows(settings_rows),
            "columns": columns, "sheets": names}


FIELD_DEFS = {
    "measurement": [
        {"key": "batch_id", "label": "Batch", "type": "text", "group_name": "Backfill", "filterable": True},
        {"key": "review_needed", "label": "Review needed", "type": "bool", "group_name": "Backfill", "show_in_table": True},
        {"key": "evidence", "label": "Evidence", "type": "json", "group_name": "Backfill"},
        {"key": "thermal_history", "label": "Thermal history", "type": "text", "group_name": "Conditions"},
        {"key": "replicate", "label": "Replicate", "type": "integer", "group_name": "Conditions"},
        {"key": "executed_at", "label": "Executed at", "type": "json", "group_name": "Session"},
        {"key": "test_name", "label": "Test name", "type": "text", "group_name": "Session"},
        {"key": "module_name", "label": "Module name", "type": "text", "group_name": "Session"},
        {"key": "clarius_version", "label": "Clarius version", "type": "text", "group_name": "Session"},
        {"key": "source_path", "label": "Source path", "type": "text", "group_name": "Backfill", "filterable": True},
        {"key": "columns", "label": "Columns", "type": "json", "group_name": "Backfill"},
    ],
    "sample": [
        {"key": "batch_id", "label": "Batch", "type": "text", "group_name": "Backfill", "filterable": True},
        {"key": "review_needed", "label": "Review needed", "type": "bool", "group_name": "Backfill", "show_in_table": True},
        {"key": "evidence", "label": "Evidence", "type": "json", "group_name": "Backfill"},
        {"key": "source_folder", "label": "Source folder", "type": "text", "group_name": "Backfill", "show_in_table": True},
    ],
}


def ensure_fields(client):
    """Create missing backfill field definitions; idempotent, prints what it created."""
    created = []
    for entity, defs in FIELD_DEFS.items():
        existing = {item["key"] for item in
                    client.get("/api/field-definitions", {"entity": entity, "include_inactive": 1})["items"]}
        for spec in defs:
            if spec["key"] in existing:
                continue
            body = {"entity": entity, "key": spec["key"], "label": spec.get("label", spec["key"]),
                    "help": None, "type": spec["type"], "options_list_key": None, "unit": None,
                    "required": False, "sort_order": 900, "group_name": spec.get("group_name", "Backfill"),
                    "active": True, "column_name": None, "show_in_table": spec.get("show_in_table", False),
                    "filterable": spec.get("filterable", False), "min": None, "max": None,
                    "regex": None, "default_value": None}
            client.post("/api/field-definitions", body)
            created.append(entity + "." + spec["key"])
    print("created field definitions: " + (", ".join(created) if created else "none (all present)"))
    return created


def upload_with_kind(client, measurement_id, path, kind, content_type):
    size = os.path.getsize(path)
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    info = client.post("/api/files/upload-url", {"measurement_id": measurement_id,
                                                 "filename": os.path.basename(path),
                                                 "size_bytes": size, "sha256": digest.hexdigest(),
                                                 "kind": kind})
    headers = {"Content-Type": content_type}
    if info.get("token"):
        headers["Authorization"] = "Bearer " + info["token"]
    with open(path, "rb") as handle:
        response = client.session.put(info["signed_url"], data=handle, headers=headers, timeout=300)
    if not response.ok:
        raise ApiError("upload_failed", response.text or response.reason)
    return client.post("/api/files/" + info["file_id"] + "/register",
                       {"size_bytes": size, "sha256": digest.hexdigest()})


def mark(dest, status, evidence, key, value, cls, source):
    if value is None:
        return
    dest[key] = value
    status[key] = "confirmed"
    evidence[key] = {"class": cls, "source": source}


def queue_item(entity, sample_id, source_path, field, candidate, reason, evidence_seen):
    return {"entity": entity, "sample_id": sample_id, "source_path": source_path, "field": field,
            "candidate_value": candidate, "reason": reason, "evidence_seen": evidence_seen}


def sample_queue_items(sample_id, folder_name):
    nothing = "not stated in the folder name, filenames, or workbook metadata"
    return [
        queue_item("sample", sample_id, "", "stack",
                   "20 nm Hf-cap stack from folder token 20Hfcap — layer unknown",
                   "only the folder token 20Hfcap is proven (E3); layer materials and thicknesses are not stated anywhere",
                   "folder name: " + folder_name),
        queue_item("sample", sample_id, "", "substrate", None, nothing, "none"),
        queue_item("sample", sample_id, "", "fab_location", None, nothing, "none"),
        queue_item("sample", sample_id, "", "fabricated_on", None, nothing, "none"),
        queue_item("sample", sample_id, "", "pad_shape", "circle",
                   "Harsh's 45 nm series used circles; not stated for this folder",
                   "folder suffix HY (E4 owner declaration); no pad shape row in workbooks"),
    ]


def find_screenshot(folder, xlsx_name):
    stem = Path(xlsx_name).stem.lower()
    while stem.endswith(".xlsx"):
        stem = stem[:-5]
    for entry in sorted(folder.iterdir()):
        if entry.suffix.lower() not in SCREENSHOT_EXTS:
            continue
        base = entry.name.lower()[: -len(entry.suffix)]
        while base.endswith(".xlsx"):
            base = base[:-5]
        if base == stem:
            return entry
    return None


def build_sample_plan(folder_name, batch):
    cols, meta, status, evidence = {}, {}, {}, {}
    mark(cols, status, evidence, "owner", "harsh_yellai", "E4", EVIDENCE_OWNER)
    mark(cols, status, evidence, "fabricated_by", "harsh_yellai", "E4", EVIDENCE_OWNER)
    mark(meta, status, evidence, "source_folder", folder_name, "E3", "folder:" + folder_name)
    mark(meta, status, evidence, "batch_id", batch, "E4", "cli --batch " + batch)
    mark(meta, status, evidence, "review_needed", True, "E4", "cli plan: unproven fields queued for review")
    body = {"sample_id": SAMPLE_ID, "label": "Harsh 20 nm high-T series (%s)" % folder_name,
            "owner": cols["owner"], "fabricated_by": cols["fabricated_by"],
            "meta": meta, "meta_status": status}
    meta["evidence"] = evidence
    status["evidence"] = "confirmed"
    return {"type": "sample", "sample_id": SAMPLE_ID, "body": body,
            "queue": sample_queue_items(SAMPLE_ID, folder_name)}


def build_measurement_plan(xlsx_path, source_path, batch, sample_id, folder_name,
                           temperature, thermal, wb, fname_info):
    queue = []
    cols, meta, status, evidence = {}, {}, {}, {}
    runs, run_numbers, columns = wb["runs"], wb["run_numbers"], wb["columns"]
    filename = Path(source_path).name

    measured_on = measured_on_from_runs(runs)
    mark(cols, status, evidence, "measured_on", measured_on, "E1", "Settings!Last Executed")
    exec_list = executed_at_list(runs)
    if exec_list:
        mark(meta, status, evidence, "executed_at", exec_list, "E1", "Settings!Last Executed")
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "measured_on", None,
                                "no Last Executed rows parsed from Settings", "Settings sheet had no usable timestamps"))
    if run_numbers:
        mark(cols, status, evidence, "run_numbers", run_numbers, "E1", "workbook sheet names")
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "run_numbers", None,
                                "no Run<N> sheets found in workbook", "sheets: " + ", ".join(wb["sheets"])))
    versions = sorted({r["clarius_version"] for r in runs if r.get("clarius_version")})
    if versions:
        mark(cols, status, evidence, "instrument", "k4200a_clarius", "E1", "Settings!Clarius+ Version")
        mark(meta, status, evidence, "clarius_version", ", ".join(versions), "E1", "Settings!Clarius+ Version")
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "instrument", None,
                                "no Clarius+ Version row in Settings", "instrument unknown"))
    test_names = sorted({r["test_name"] for r in runs if r.get("test_name")})
    module_names = sorted({r["module_name"] for r in runs if r.get("module_name")})
    token_kind = fname_info["kind"] if fname_info else None
    kind, kind_cls, kind_src, kind_reason = resolve_kind(token_kind, test_names, module_names)
    if kind:
        mark(cols, status, evidence, "kind", kind, kind_cls, kind_src)
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "kind",
                                token_kind or (sorted(test_names + module_names) or None),
                                kind_reason or "kind could not be proven", "Settings!Test Name/Module Name + filename"))
    if test_names:
        mark(meta, status, evidence, "test_name", ", ".join(test_names), "E1", "Settings!Test Name")
    if module_names:
        mark(meta, status, evidence, "module_name", ", ".join(module_names), "E1", "Settings!Module Name")

    if fname_info:
        mark(cols, status, evidence, "pad_dim_um", fname_info["pad_dim_um"], "E2", "filename:" + filename)
        mark(meta, status, evidence, "replicate", fname_info["replicate"], "E2", "filename:" + filename)
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "pad_dim_um", None,
                                "filename does not match <pad>-<KIND>-<rep>.xlsx", "filename: " + filename))
    if columns:
        mark(meta, status, evidence, "columns", columns, "E1", "data sheet header row")

    sweeps = {(r["sweep_v"], r["sweep_src"]) for r in runs if r.get("sweep_v") is not None}
    if len(sweeps) == 1:
        value, src = next(iter(sweeps))
        mark(meta, status, evidence, "sweep_v", value, "E1", "Settings!" + src)
    freqs = {(r["frequency_khz"], r["freq_src"]) for r in runs if r.get("frequency_khz") is not None}
    if len(freqs) == 1:
        value, src = next(iter(freqs))
        mark(meta, status, evidence, "frequency_khz", value, "E1", "Settings!" + src)
    settings_keys = sorted({k for r in runs for k in r["keys"]})
    settings_seen = "Settings rows seen: " + ", ".join(settings_keys[:12])
    if len(sweeps) != 1:
        queue.append(queue_item("measurement", sample_id, source_path, "sweep_v", None,
                                "no Start/Stop sweep row with a number in Settings" if not sweeps
                                else "sweep values disagree across runs", settings_seen))
    if len(freqs) != 1:
        queue.append(queue_item("measurement", sample_id, source_path, "frequency_khz", None,
                                "no frequency row with a stated unit in Settings" if not freqs
                                else "frequencies disagree across runs", settings_seen))

    if temperature is not None:
        mark(cols, status, evidence, "temperature_c", temperature, "E3", "folder:" + folder_name)
    else:
        queue.append(queue_item("measurement", sample_id, source_path, "temperature_c", None,
                                "subfolder name has no temperature token", "folder: " + folder_name))
    mark(meta, status, evidence, "thermal_history", thermal, "E3", "folder:" + folder_name)
    mark(cols, status, evidence, "measured_by", "harsh_yellai", "E4", EVIDENCE_OWNER)
    mark(meta, status, evidence, "source_path", source_path, "E1", "workbook path")
    mark(meta, status, evidence, "batch_id", batch, "E4", "cli --batch " + batch)
    mark(meta, status, evidence, "review_needed", True, "E4", "cli plan: unproven fields queued for review")
    queue.append(queue_item("measurement", sample_id, source_path, "pad_shape", "circle",
                            "Harsh's 45 nm series used circles; not stated for this folder",
                            "folder suffix HY (E4 owner declaration); no pad shape row in workbooks"))
    queue.append(queue_item("measurement", sample_id, source_path, "probe_station", "hot_chuck_station",
                            "guess: temperatures reach 600 C, consistent with a hot chuck; never stated",
                            "folder " + folder_name + " (E3 temperature token only)"))

    files = [{"path": str(xlsx_path), "kind": "raw_xls", "content_type": "application/octet-stream"}]
    screenshot = find_screenshot(xlsx_path.parent, xlsx_path.name)
    if screenshot:
        files.append({"path": str(screenshot), "kind": "plot_png",
                      "content_type": SCREENSHOT_EXTS[screenshot.suffix.lower()]})
    body = dict(cols)
    body["meta"] = meta
    body["meta_status"] = status
    meta["evidence"] = evidence
    status["evidence"] = "confirmed"
    return {"type": "measurement", "sample_id": sample_id, "source_path": source_path,
            "body": body, "files": files, "queue": queue}


def plan(folder, batch):
    folder = Path(folder).resolve()
    if not folder.is_dir():
        sys.exit("folder not found: %s" % folder)
    data_root = folder.parent
    sample_plan = build_sample_plan(folder.name, batch)
    measurements, failures = [], []
    for sub in sorted(p for p in folder.iterdir() if p.is_dir()):
        temperature, thermal = parse_folder(sub.name)
        for xlsx in sorted(sub.glob("*.xlsx")):
            source_path = xlsx.relative_to(data_root).as_posix()
            try:
                wb = read_workbook(xlsx)
            except Exception as exc:
                failures.append((source_path, "%s: %s" % (type(exc).__name__, exc)))
                continue
            measurements.append(build_measurement_plan(
                xlsx, source_path, batch, sample_plan["sample_id"], sub.name,
                temperature, thermal, wb, parse_filename(xlsx.name)))
    return sample_plan, measurements, failures


def print_report(sample_plan, measurements, failures):
    print("sample: %s | %s" % (sample_plan["sample_id"], sample_plan["body"]["label"]))
    print("workbooks planned: %d; failed to parse: %d" % (len(measurements), len(failures)))
    for source_path, error in failures:
        print("  FAILED %s: %s" % (source_path, error))
    kinds = Counter(m["body"].get("kind") or "<queued>" for m in measurements)
    print("per kind: " + ", ".join("%s=%d" % kv for kv in sorted(kinds.items())))
    temps = Counter(str(m["body"].get("temperature_c")) for m in measurements)
    print("per temperature: " + ", ".join("%s=%d" % kv for kv in sorted(temps.items())))
    fields = Counter()
    for m in measurements:
        for key, ev in m["body"]["meta"]["evidence"].items():
            fields[(key, ev["class"])] += 1
    print("fields set (field/class=count):")
    for (key, cls), count in sorted(fields.items()):
        print("  %-16s %-3s %d" % (key, cls, count))
    reasons = Counter((q["field"], q["reason"]) for m in measurements for q in m["queue"])
    print("measurement queue items by reason:")
    for (field, reason), count in sorted(reasons.items()):
        print("  %-16s x%d  %s" % (field, count, reason))
    paired = sum(1 for m in measurements if len(m["files"]) > 1)
    unpaired = [m["source_path"] for m in measurements if len(m["files"]) == 1]
    total_bytes = sum(os.path.getsize(f["path"]) for m in measurements for f in m["files"])
    print("screenshots paired: %d/%d%s" % (paired, len(measurements),
                                           ("; unpaired: " + ", ".join(unpaired)) if unpaired else ""))
    print("total bytes to upload: %d (%.1f MB)" % (total_bytes, total_bytes / 1e6))
    print("sample queue items: %s" % ", ".join(q["field"] for q in sample_plan["queue"]))


def read_plan(path):
    lines = []
    with open(path, "r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                lines.append(json.loads(line))
    return lines


def cmd_plan(args):
    sample_plan, measurements, failures = plan(args.folder, args.batch)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(sample_plan, ensure_ascii=False) + "\n")
        for m in measurements:
            handle.write(json.dumps(m, ensure_ascii=False) + "\n")
    print_report(sample_plan, measurements, failures)
    print("wrote %s" % out)


def cmd_queue(args):
    lines = read_plan(args.plan)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    rows = 0
    with out.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["entity", "sample_id", "source_path", "field", "candidate_value",
                         "reason", "evidence_seen"])
        for line in lines:
            for item in line["queue"]:
                writer.writerow([item["entity"], item["sample_id"], item.get("source_path", ""),
                                 item["field"],
                                 "" if item["candidate_value"] is None else item["candidate_value"],
                                 item["reason"], item.get("evidence_seen", "")])
                rows += 1
    print("wrote %s (%d queue rows)" % (out, rows))


def save_state(path, state):
    path.write_text(json.dumps(state, indent=1), encoding="utf-8")


def cmd_run(args):
    client = make_client()
    ensure_fields(client)
    lines = read_plan(args.plan)
    warnings_all = []
    for sample_plan in [l for l in lines if l["type"] == "sample"]:
        sid = sample_plan["sample_id"]
        try:
            client.get("/api/samples/" + sid)
            print("sample %s exists; reusing" % sid)
        except ApiError as exc:
            if exc.code != "not_found":
                raise
            result = client.post("/api/samples", sample_plan["body"])
            warnings_all += result.get("warnings", [])
            print("created sample %s" % sid)
    measurement_lines = [l for l in lines if l["type"] == "measurement"]
    state_path = Path(args.plan).parent / "state.json"
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {"done": {}}
    existing_cache, created, skipped, errors, uploaded = {}, 0, 0, 0, 0
    for mp in measurement_lines:
        if args.limit is not None and created >= args.limit:
            break
        sid, source_path = mp["sample_id"], mp["source_path"]
        if source_path in state["done"]:
            skipped += 1
            continue
        if sid not in existing_cache:
            items = client.get("/api/samples/%s/measurements" % sid, {"limit": 200})["items"]
            existing_cache[sid] = {(m.get("meta") or {}).get("source_path") for m in items}
        if source_path in existing_cache[sid]:
            state["done"][source_path] = "exists"
            save_state(state_path, state)
            skipped += 1
            continue
        try:
            result = client.post("/api/samples/%s/measurements" % sid, mp["body"])
        except ApiError as exc:
            print("ERROR %s: %s: %s" % (source_path, exc.code, exc.message))
            errors += 1
            continue
        measurement_id = result["measurement"]["id"]
        warnings_all += result.get("warnings", [])
        wanted = mp["files"] if not args.skip_screenshots else \
            [f for f in mp["files"] if f["kind"] == "raw_xls"]
        for entry in wanted:
            try:
                upload_with_kind(client, measurement_id, entry["path"], entry["kind"], entry["content_type"])
                uploaded += os.path.getsize(entry["path"])
            except ApiError as exc:
                if exc.code == "duplicate_file":
                    print("SKIP duplicate file %s" % entry["path"])
                else:
                    print("ERROR upload %s: %s: %s" % (entry["path"], exc.code, exc.message))
                    errors += 1
        state["done"][source_path] = measurement_id
        save_state(state_path, state)
        created += 1
        time.sleep(0.2)
    print("tally: created=%d skipped=%d errors=%d bytes_uploaded=%d" % (created, skipped, errors, uploaded))
    if warnings_all:
        print("API warnings (%d):" % len(warnings_all))
        for warning in sorted(set(warnings_all)):
            print("  " + warning)


def cmd_rollback(args):
    if not args.yes:
        sys.exit("rollback requires --yes")
    client = make_client()
    samples = client.get("/api/samples", {"meta.batch_id": args.batch, "limit": 200})["items"]
    deleted_m, deleted_s = 0, 0
    for sample in samples:
        sid = sample["sample_id"]
        items = client.get("/api/samples/%s/measurements" % sid, {"meta.batch_id": args.batch, "limit": 200})["items"]
        for m in [x for x in items if (x.get("meta") or {}).get("batch_id") == args.batch]:
            client.delete("/api/measurements/" + m["id"])
            deleted_m += 1
            print("deleted measurement %s (%s)" % (m["id"], sid))
        rest = client.get("/api/samples/%s/measurements" % sid, {"limit": 1})
        if not rest["total"]:
            client.delete("/api/samples/" + sid)
            deleted_s += 1
            print("deleted sample %s (no measurements left)" % sid)
    print("rollback done: %d measurements, %d samples deleted" % (deleted_m, deleted_s))


def main(argv=None):
    parser = argparse.ArgumentParser(prog="backfill.py",
                                     description="Sure-only backfill for the Agni Data Vault")
    sub = parser.add_subparsers(dest="command", required=True)
    plan_p = sub.add_parser("plan")
    plan_p.add_argument("--folder", required=True)
    plan_p.add_argument("--batch", required=True)
    plan_p.add_argument("--out", default="backfill/plan.jsonl")
    run_p = sub.add_parser("run")
    run_p.add_argument("--plan", default="backfill/plan.jsonl")
    run_p.add_argument("--limit", type=int)
    run_p.add_argument("--skip-screenshots", action="store_true")
    queue_p = sub.add_parser("queue")
    queue_p.add_argument("--plan", default="backfill/plan.jsonl")
    queue_p.add_argument("--out", default="backfill/review_queue.csv")
    back_p = sub.add_parser("rollback")
    back_p.add_argument("--batch", required=True)
    back_p.add_argument("--yes", action="store_true")
    args = parser.parse_args(argv)
    {"plan": cmd_plan, "run": cmd_run, "queue": cmd_queue, "rollback": cmd_rollback}[args.command](args)
    return 0


if __name__ == "__main__":
    sys.exit(main())
